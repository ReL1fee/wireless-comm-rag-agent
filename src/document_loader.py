"""
文档加载与分块模块
支持 Markdown、PDF、Word(docx) 格式
"""
import os
from langchain_core.documents import Document
from langchain_community.document_loaders import DirectoryLoader, TextLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from src.config import DATA_DIR, CHUNK_SIZE, CHUNK_OVERLAP


def _get_splitter():
    return RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        separators=["\n## ", "\n### ", "\n", "。", ".", " "],
        length_function=len,
    )


def load_markdown_files(data_dir: str = DATA_DIR) -> list:
    """从目录加载所有 markdown 文件"""
    loader = DirectoryLoader(
        data_dir,
        glob="**/*.md",
        loader_cls=TextLoader,
        loader_kwargs={"encoding": "utf-8"},
        show_progress=True,
    )
    return loader.load()


def load_pdf(filepath: str) -> list[Document]:
    """加载单个 PDF 文件，返回 Document 列表"""
    try:
        from pypdf import PdfReader
        reader = PdfReader(filepath)
        text_parts = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text and page_text.strip():
                text_parts.append(page_text)

        full_text = "\n\n".join(text_parts)
        if not full_text.strip():
            raise ValueError("PDF 文件中未提取到文字内容（仅支持文字PDF）")

        filename = os.path.basename(filepath)
        return [Document(page_content=full_text, metadata={"source": filename})]
    except ImportError:
        raise ImportError("请先安装 pypdf: pip install pypdf")


def load_docx(filepath: str) -> list[Document]:
    """加载单个 Word(docx) 文件，返回 Document 列表"""
    try:
        from docx import Document as DocxDocument
        doc = DocxDocument(filepath)
        text_parts = [para.text for para in doc.paragraphs if para.text.strip()]
        full_text = "\n".join(text_parts)

        if not full_text.strip():
            raise ValueError("Word 文件中未提取到文字内容")

        filename = os.path.basename(filepath)
        return [Document(page_content=full_text, metadata={"source": filename})]
    except ImportError:
        raise ImportError("请先安装 python-docx: pip install python-docx")


def load_uploaded_file(filepath: str) -> list[Document]:
    """根据文件类型自动选择加载方式，返回 Document 列表"""
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".pdf":
        return load_pdf(filepath)
    elif ext in (".docx", ".doc"):
        return load_docx(filepath)
    elif ext == ".md":
        loader = TextLoader(filepath, encoding="utf-8")
        return loader.load()
    elif ext == ".txt":
        loader = TextLoader(filepath, encoding="utf-8")
        return loader.load()
    else:
        raise ValueError(f"不支持的文件格式: {ext}，目前支持 PDF、Word、Markdown、TXT")


def split_documents(documents: list) -> list:
    """将文档分割为适合检索的文本块"""
    chunks = _get_splitter().split_documents(documents)
    print(f"[文档分块] 共分割为 {len(chunks)} 个文本块 (块大小={CHUNK_SIZE}，重叠={CHUNK_OVERLAP})")
    return chunks


def get_chunks(data_dir: str = DATA_DIR) -> list:
    """从 data/ 目录加载所有支持的文档并分块"""
    documents = []

    # 加载 MD 文件
    md_docs = load_markdown_files(data_dir)
    documents.extend(md_docs)

    # 加载 PDF 文件
    for root, _, files in os.walk(data_dir):
        for f in files:
            if f.lower().endswith(".pdf"):
                filepath = os.path.join(root, f)
                try:
                    documents.extend(load_pdf(filepath))
                    print(f"[文档加载] PDF: {f}")
                except Exception as e:
                    print(f"[警告] 跳过 {f}: {e}")

    # 加载 DOCX 文件
    for root, _, files in os.walk(data_dir):
        for f in files:
            if f.lower().endswith(".docx"):
                filepath = os.path.join(root, f)
                try:
                    documents.extend(load_docx(filepath))
                    print(f"[文档加载] DOCX: {f}")
                except Exception as e:
                    print(f"[警告] 跳过 {f}: {e}")

    if not documents:
        raise RuntimeError(f"未在 {data_dir} 中找到任何文档，请先添加知识库文件")

    print(f"[文档加载] 共加载 {len(documents)} 个文档")
    return split_documents(documents)


def process_uploaded_file(filepath: str) -> list:
    """处理用户上传的单个文件：加载 + 分块，返回 chunks"""
    documents = load_uploaded_file(filepath)
    if not documents:
        raise RuntimeError(f"文件内容为空: {filepath}")
    print(f"[上传处理] {os.path.basename(filepath)} → {len(documents)} 个文档")
    return split_documents(documents)
